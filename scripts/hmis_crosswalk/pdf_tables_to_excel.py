"""
Pull every table out of a PDF (default: the HUD HMIS Data Dictionary) into
an Excel workbook.

Usage:
    python scripts/hmis_crosswalk/pdf_tables_to_excel.py
    python scripts/hmis_crosswalk/pdf_tables_to_excel.py path/or/url/to/file.pdf -o out.xlsx
    python scripts/hmis_crosswalk/pdf_tables_to_excel.py --no-merge --pages 10-40
    python scripts/hmis_crosswalk/pdf_tables_to_excel.py --union 1,4,10-20   # or --union all

    # HMIS CSV Format Specifications: one table per CSV file, headed "Client.csv" or "Custom*.csv"
    python scripts/hmis_crosswalk/pdf_tables_to_excel.py HMIS-CSV-Format-Specifications.pdf \
        -o csv_fields.xlsx --heading-regex '^[\\w*]+\\.csv$' --union-header Name --compact

Output workbook:
    All Tables    union of the selected tables: data rows stacked, columns
                  aligned by header name, tagged with source sheet, heading and
                  pages. By default only tables whose first data row starts with
                  "Element Name" are included; pick others with --union, or by
                  a header cell with --union-header.
    Index        one row per table: sheet name, source pages, heading, size
    T001, T002... one sheet per table, preceded by its heading and pages

Tables that run across a page break (same column count, starting at the top
of the next page) are stitched into one table, and a repeated header row on
the continuation page is dropped. Each table is labelled with the nearest
data-element heading above it (e.g. "3.01 Name", "P4 Connection with SOAR"),
carried forward from earlier pages when the page itself has none (change the
pattern with --heading-regex). Text inside tables is never taken as a heading,
and tables nested inside another table's cell are skipped.

With --compact, columns that pdfplumber splits into sub-columns (common with
doubled cell borders, as in the HMIS CSV spec) are folded back into the header
column beside them.
"""
import argparse
import re
import sys
import tempfile
import urllib.request
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_URL = "https://files.hudexchange.info/resources/documents/HMIS-Data-Dictionary.pdf"

# Data-element headings in the HMIS dictionary: "2.02 Project Information",
# "3.917 Prior Living Situation", "4.12 Current Living Situation",
# "P4 Connection with SOAR", "R7 General Health Status", "V2 ...", "W5 ...", "C2 ..."
HEADING_RE = re.compile(r"^(\d+\.\d+[A-Z]?|[A-Z]\d{1,2}[A-Z]?)\s+[A-Z][\w&/(),'\- ]{2,}$")

# Where a table must start / end (as a fraction of page height) to count as
# continuing across a page break.
TOP_FRACTION = 0.20
BOTTOM_FRACTION = 0.80

# Default --union filter: first cell of the first non-header row.
UNION_FIRST_CELL = "Element Name"

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def fetch_pdf(source):
    """Return a local path for `source`, downloading it if it is a URL."""
    if not re.match(r"https?://", source):
        return Path(source)
    dest = Path(tempfile.gettempdir()) / Path(source).name
    if not dest.exists():
        print(f"Downloading {source} ...", file=sys.stderr)
        req = urllib.request.Request(source, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req) as resp, open(dest, "wb") as fh:
            fh.write(resp.read())
    return dest


def parse_pages(spec, n_pages):
    """'1-5,8' -> [0, 1, 2, 3, 4, 7] (0-based)."""
    if not spec:
        return list(range(n_pages))
    pages = []
    for part in spec.split(","):
        lo, _, hi = part.partition("-")
        pages.extend(range(int(lo) - 1, int(hi or lo)))
    return [p for p in pages if 0 <= p < n_pages]


def clean_cell(value):
    if value is None:
        return ""
    # Collapse the hard line wraps pdfplumber keeps inside cells.
    return re.sub(r"\s*\n\s*", " ", str(value)).strip()


def clean_rows(rows):
    rows = [[clean_cell(c) for c in row] for row in rows]
    return [r for r in rows if any(r)]


def inside(inner, outer, tol=1):
    return (inner[0] >= outer[0] - tol and inner[1] >= outer[1] - tol
            and inner[2] <= outer[2] + tol and inner[3] <= outer[3] + tol)


def find_heading(page, table_top, table_boxes, heading_re):
    """Closest heading above `table_top` on this page, ignoring text inside tables."""
    lines = page.extract_text_lines(strip=True) if hasattr(page, "extract_text_lines") else []
    heading = None
    for line in lines:
        if line["bottom"] > table_top:
            break
        box = (line["x0"], line["top"], line["x1"], line["bottom"])
        if any(inside(box, tb) for tb in table_boxes):
            continue
        text = line["text"].strip()
        if heading_re.match(text):
            heading = text
    return heading


def compact_columns(rows):
    """Fold sub-columns into the header column beside them.

    With doubled cell borders pdfplumber can split one visual column into
    several, leaving the header text in one sub-column and the data one to its
    left. Each column with data but no header joins an adjacent header column
    (the right one first) unless that would put two different values in one
    cell; identical text is fine, since some PDFs draw a cell's text twice.
    """
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    header = rows[0]
    groups = {j: [j] for j in range(n) if header[j]}

    def text(r, cols):
        return " ".join(dict.fromkeys(r[k] for k in cols if r[k]))

    for j in range(n):
        if header[j] or not any(r[j] for r in rows):
            continue
        for a in (j + 1, j - 1):
            if a in groups and not any(r[j] and text(r, groups[a]) not in ("", r[j]) for r in rows):
                groups[a].append(j)
                break
        else:
            groups[j] = [j]  # keep as its own column
    order = sorted(groups)
    return [[text(r, groups[a]) for a in order] for r in rows]


def extract_tables(pdf_path, page_spec=None, merge=True, heading_re=HEADING_RE, compact=False):
    tables = []
    last_heading = ""
    with pdfplumber.open(pdf_path) as pdf:
        for pno in parse_pages(page_spec, len(pdf.pages)):
            page = pdf.pages[pno]
            found = page.find_tables()
            # Drop tables nested in another table's cell (e.g. a boxed regex in a Notes cell).
            found = [t for t in found
                     if not any(o is not t and inside(t.bbox, o.bbox) for o in found)]
            found.sort(key=lambda t: t.bbox[1])
            boxes = [t.bbox for t in found]
            if not found:
                # Still track headings so the next table gets the right label.
                h = find_heading(page, page.height, boxes, heading_re)
                last_heading = h or last_heading
                continue
            for tbl in found:
                rows = clean_rows(tbl.extract())
                if not rows:
                    continue
                top, bottom = tbl.bbox[1], tbl.bbox[3]
                heading = find_heading(page, top, boxes, heading_re)
                prev = tables[-1] if tables else None
                is_continuation = (
                    merge
                    and prev is not None
                    and heading is None
                    and prev["last_page"] == pno - 1
                    and prev["ends_low"]
                    and top < page.height * TOP_FRACTION
                    and len(rows[0]) == len(prev["rows"][0])
                )
                if is_continuation:
                    if rows[0] == prev["rows"][0]:
                        rows = rows[1:]  # repeated header
                    prev["rows"].extend(rows)
                    prev["last_page"] = pno
                    prev["ends_low"] = bottom > page.height * BOTTOM_FRACTION
                    continue
                last_heading = heading or last_heading
                tables.append({
                    "heading": last_heading,
                    "first_page": pno,
                    "last_page": pno,
                    "ends_low": bottom > page.height * BOTTOM_FRACTION,
                    "rows": rows,
                })
            # Headings below the last table on the page apply to the next page.
            h = find_heading(page, page.height, boxes, heading_re)
            last_heading = h or last_heading
    if compact:
        for t in tables:
            t["rows"] = compact_columns(t["rows"])
    return tables


def autosize(ws, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), max_width)


def table_headers(header_row):
    """Header names for one table: blanks become 'Column N', repeats get a suffix."""
    names, seen = [], {}
    for i, h in enumerate(header_row, 1):
        name = h or f"Column {i}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return names


def select_union(tables, spec=None, header=None):
    """0-based indices of the tables to union.

    spec: None for the default filter (first data cell == UNION_FIRST_CELL),
    "all", or a comma list of table numbers / ranges, e.g. "1,4,10-20" or "T001-T005".
    header: instead, pick tables whose header row has a cell equal to this text.
    """
    if header is not None:
        want = header.casefold()
        return [i for i, t in enumerate(tables)
                if any(c.casefold() == want for c in t["rows"][0])]
    if spec is None:
        want = UNION_FIRST_CELL.casefold()
        return [i for i, t in enumerate(tables)
                if len(t["rows"]) > 1 and t["rows"][1][0].casefold() == want]
    if spec.strip().lower() == "all":
        return list(range(len(tables)))
    picked = set()
    for part in spec.split(","):
        lo, _, hi = part.strip().upper().replace("T", "").partition("-")
        picked.update(range(int(lo), int(hi or lo) + 1))
    return [n - 1 for n in sorted(picked) if 1 <= n <= len(tables)]


def write_union(ws, tables, sheet_names):
    """Stack every table's data rows into one sheet, aligning columns by header name."""
    columns = []
    for t in tables:
        prev = None
        for name in table_headers(t["rows"][0]):
            if name not in columns:
                # Keep a new column next to its neighbour in this table's own order.
                columns.insert(columns.index(prev) + 1 if prev else 0, name)
            prev = name
    lead = ["Table", "Heading", "Pages"]
    ws.append(lead + columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    for t, name in zip(tables, sheet_names):
        headers = table_headers(t["rows"][0])
        for row in t["rows"][1:]:
            values = dict(zip(headers, row))
            ws.append([name, t["heading"], t["pages"]] + [values.get(c, "") for c in columns])

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)


def write_workbook(tables, out_path, union_idx):
    for t in tables:
        t["pages"] = (f"{t['first_page'] + 1}" if t["first_page"] == t["last_page"]
                      else f"{t['first_page'] + 1}-{t['last_page'] + 1}")
    sheet_names = [f"T{i:03d}" for i in range(1, len(tables) + 1)]

    wb = Workbook()
    write_union(wb.active, [tables[i] for i in sorted(union_idx)], [sheet_names[i] for i in sorted(union_idx)])
    wb.active.title = "All Tables"
    index = wb.create_sheet("Index")
    index.append(["Sheet", "Heading", "Pages", "Rows", "Columns", "In All Tables"])
    for cell in index[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    for i, (t, name) in enumerate(zip(tables, sheet_names)):
        pages = t["pages"]
        n_cols = max(len(r) for r in t["rows"])

        ws = wb.create_sheet(name)
        ws.append([t["heading"] or "(no heading found)"])
        ws.append([f"PDF page(s) {pages}"])
        ws.append([])
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"].font = Font(italic=True)
        for r in t["rows"]:
            ws.append(r)
        for cell in ws[4]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A5"
        autosize(ws)

        index.append([name, t["heading"], pages, len(t["rows"]), n_cols,
                      "Yes" if i in union_idx else "No"])
        index.cell(row=index.max_row, column=1).hyperlink = f"#'{name}'!A1"
        index.cell(row=index.max_row, column=1).style = "Hyperlink"

    index.freeze_panes = "A2"
    autosize(index, max_width=80)
    wb.save(out_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", nargs="?", default=DEFAULT_URL, help="PDF path or URL")
    ap.add_argument("-o", "--output", default="HMIS-Data-Dictionary-tables.xlsx")
    ap.add_argument("--pages", help="1-based page ranges to scan, e.g. '1-20,35'")
    ap.add_argument("--no-merge", action="store_true", help="don't stitch tables across page breaks")
    ap.add_argument("--union", metavar="TABLES",
                    help="tables to include in the 'All Tables' sheet: 'all' or numbers/ranges "
                         "like '1,4,10-20' (default: tables whose first data row starts with "
                         f"'{UNION_FIRST_CELL}')")
    ap.add_argument("--union-header", metavar="TEXT",
                    help="union the tables whose header row has a cell equal to TEXT, e.g. 'Name'")
    ap.add_argument("--heading-regex", default=HEADING_RE.pattern,
                    help="pattern for the heading lines that label tables "
                         "(default: HMIS data-element headings like '3.01 Name')")
    ap.add_argument("--compact", action="store_true",
                    help="fold sub-columns pdfplumber split off back into their header column")
    args = ap.parse_args(argv)

    pdf_path = fetch_pdf(args.source)
    tables = extract_tables(pdf_path, args.pages, merge=not args.no_merge,
                            heading_re=re.compile(args.heading_regex), compact=args.compact)
    if not tables:
        sys.exit("No tables found.")
    union_idx = select_union(tables, args.union, args.union_header)
    if not union_idx:
        print("Warning: no tables selected for the 'All Tables' sheet.", file=sys.stderr)
    write_workbook(tables, args.output, set(union_idx))
    print(f"Wrote {len(tables)} tables to {args.output} ({len(union_idx)} in 'All Tables')")


if __name__ == "__main__":
    main()
