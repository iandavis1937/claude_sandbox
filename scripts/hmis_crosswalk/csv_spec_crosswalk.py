"""
Build a crosswalk from the HMIS CSV Format Specifications to HMIS Data
Dictionary element/field numbers, from the spec's own text, for the columns
its DE# column does not cover. Feeds join_dd_csv.py --crosswalk.

Usage:
    python scripts/hmis_crosswalk/csv_spec_crosswalk.py HMIS-CSV-Format-Specifications.pdf HMIS_CSV_fields.xlsx
    python scripts/hmis_crosswalk/csv_spec_crosswalk.py spec.pdf csv_fields.xlsx -o HMIS_CSV_crosswalk.xlsx

Inputs: the spec PDF, and its field list from
    pdf_tables_to_excel.py spec.pdf --heading-regex '^[\\w*]+\\.csv$' --union-header Name --compact

Output workbook:
    File Elements  the data elements each CSV file includes, from the spec's
                   "Disabilities.csv includes data from: • 4.05 Physical
                   Disability (DisabilityType 5) ..." lists and "Project.csv
                   includes all the fields from data element 2.02 ..." sentences
    Field Map      CSV column -> DD element or field, with the condition under
                   which it holds and the rule that produced it:
        discriminator  DisabilityType / RecordType <-> each listed element
                       ("DisabilityType = 5" <-> 4.05)
        date column    InformationDate / DateProvided <-> each listed element's
                       date field (role 'date')
        value column   the column whose Notes read "For <discriminator> ..."
                       (DisabilityResponse, TypeProvided) <-> each element's
                       value field; "• 141 – list P1.2" gives the exact field
        list column    a column with no DE# whose List cites a DD field (5.03.1)
        metadata       DateCreated, UserID, ... (MANUAL_COLUMNS)
        dependent      a column holding each element's Dependent X
                       (MANUAL_DEPENDENTS: Disabilities.csv IndefiniteAndImpairs)

Element-level rows (Level = element) name a role - date, value, or
discriminator - that the join resolves to one of the element's fields.
"""
import argparse
import re

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")

EL = r"\d\.\d{2,3}|[A-Z]\d{1,2}"
CSV_HEADING_RE = re.compile(r"^[\w*]+\.csv$")
# "• 4.05 Physical Disability (DisabilityType 5)", "• V8 HUD-VASH ... (RecordType 210)9"
BULLET_RE = re.compile(rf"^•\s*({EL})\s+([A-Z].*?)\s*(?:\((\w+)\s+(\d+)\))?\s*\d*$")
# "... includes all the fields from data element 2.01 Organization Information."
INCLUDES_RE = re.compile(rf"\b({EL})(?:\.[A-Z])?\s+([A-Z][\w/&’'– -]*?)(?=[.,;(]|\s+(?:about|There|which)\b|$)")
# Notes: "For RecordType • 141 – list P1.2 • 142 – list R14.2 ..."
FOR_RE = re.compile(r"\bFor (\w+)\b")
LIST_ITEM_RE = re.compile(rf"•\s*(\d+)\b[^•]*?\blist\s+((?:{EL})(?:\.\w+)?)")
# A DD field cited in the List column: "5.03.1"
LIST_FIELD_RE = re.compile(rf"^(?:{EL})\.\w+$")

DATE_COLUMNS = ("InformationDate", "DateProvided")

# Columns with no DE# that carry DD metadata fields in every file.
MANUAL_COLUMNS = {
    "DateCreated": ["5.01.1"],
    "DateUpdated": ["5.02.1"],
    "InformationDate": ["5.04.1"],
    "EnrollmentID": ["5.06.1"],
    "UserID": ["5.07.01"],  # the spec prints 5.07.1 as "5.07.01" (User.csv)
    "PersonalID": ["5.08.1"],
    "HouseholdID": ["5.09.1"],
    "ImplementationID": ["5.10.1"],
    "ProjectID": ["2.02.1", "5.05.1"],
}

# Columns of discriminated files that hold one dependent of each listed element,
# which the spec's text does not spell out: (file, column) -> dependent letter.
MANUAL_DEPENDENTS = {
    ("Disabilities.csv", "IndefiniteAndImpairs"): "A",  # "...long-continued and indefinite duration..."
}

FIELD_MAP_COLS = ["CSV File", "CSV Field", "Condition", "DE#", "Level", "Role", "Rule", "Evidence"]


def text_lines(pdf_path):
    """(page, text) for every line outside a table, in reading order."""
    with pdfplumber.open(pdf_path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            boxes = [t.bbox for t in page.find_tables()]
            for line in page.extract_text_lines():
                if any(line["top"] >= b[1] - 1 and line["bottom"] <= b[3] + 1 for b in boxes):
                    continue
                yield pno, line["text"].strip()


def file_elements(pdf_path):
    rows, seen, csv_file = [], set(), None
    for pno, line in text_lines(pdf_path):
        if CSV_HEADING_RE.match(line):
            csv_file = line
            continue
        if not csv_file:
            continue
        found = []
        m = BULLET_RE.match(line)
        if m:
            found.append(m.groups())
        elif "includes" in line and line.startswith(csv_file.split(".")[0]):
            found += [(el, name.strip(), None, None) for el, name in INCLUDES_RE.findall(line)]
        for el, name, disc_col, disc_val in found:
            if (csv_file, el) in seen:
                continue
            seen.add((csv_file, el))
            rows.append({"CSV File": csv_file, "Element": el, "Element Name": name,
                         "Discriminator": disc_col, "Value": int(disc_val) if disc_val else None,
                         "Page": pno, "Source Text": line})
    return rows


def read_fields(path):
    ws = load_workbook(path, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    return [dict(zip(header, r)) for r in it if any(r)]


def field_map(elements, fields):
    out = []
    by_file = {}
    for f in fields:
        by_file.setdefault(f["Heading"], []).append(f)

    def add(f, cond, key, level, role, rule, evidence):
        out.append({"CSV File": f["Heading"], "CSV Field": f["Name"], "Condition": cond,
                    "DE#": key, "Level": level, "Role": role, "Rule": rule, "Evidence": evidence})

    def no_de(f):
        return f.get("DE#") in (None, "", "*")

    for csv_file, cols in by_file.items():
        listed = [e for e in elements if e["CSV File"] == csv_file]
        date_col = next((f for f in cols if f["Name"] in DATE_COLUMNS and no_de(f)), None)

        # Value columns: Notes "For <discriminator> • <value> – list <key> ...".
        # Their field keys also tell the discriminator which field it marks.
        value_key = {}  # (discriminator, value) -> DD field key
        for f in cols:
            notes = str(f.get("Notes") or "")
            m = FOR_RE.search(notes)
            if not (m and no_de(f)):
                continue
            disc = m.group(1)
            explicit = {int(v): key for v, key in LIST_ITEM_RE.findall(notes)}
            for e in listed:
                if e["Discriminator"] != disc:
                    continue
                cond = f"{disc} = {e['Value']}"
                key = explicit.get(e["Value"])
                if key and key.startswith(e["Element"] + ".") and key != e["Element"]:
                    value_key[(disc, e["Value"])] = key
                    add(f, cond, key, "field", "value", "value column (Notes list)", notes)
                else:
                    add(f, cond, e["Element"], "element", "value", "value column", notes)

        for e in listed:
            el, disc, val = e["Element"], e["Discriminator"], e["Value"]
            cond = f"{disc} = {val}" if disc else ""
            if disc:
                disc_col = next((f for f in cols if f["Name"] == disc), None)
                if disc_col:
                    key = value_key.get((disc, val))
                    add(disc_col, cond, key or el, "field" if key else "element", "discriminator",
                        "discriminator", e["Source Text"])
            if date_col:
                add(date_col, cond, el, "element", "date", "date column", e["Source Text"])
            for f in cols:
                letter = MANUAL_DEPENDENTS.get((csv_file, f["Name"]))
                if letter and disc:
                    add(f, cond, f"{el}.{letter}", "field", "dependent", "dependent (manual)", "")

        for f in cols:
            if not no_de(f):
                continue
            lst = str(f.get("List") or "")
            if LIST_FIELD_RE.match(lst):
                add(f, "", lst, "field", "field", "list column", f"List = {lst}")
            for key in MANUAL_COLUMNS.get(f["Name"], []):
                add(f, "", key, "field", "field", "metadata (manual)", "")
    return out


def write_sheet(ws, columns, rows):
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf", help="HMIS CSV Format Specifications PDF")
    ap.add_argument("csv_fields", help="pdf_tables_to_excel.py output for that PDF")
    ap.add_argument("-o", "--output", default="HMIS_CSV_crosswalk.xlsx")
    args = ap.parse_args(argv)

    elements = file_elements(args.pdf)
    fmap = field_map(elements, read_fields(args.csv_fields))

    wb = Workbook()
    write_sheet(wb.active, ["CSV File", "Element", "Element Name", "Discriminator", "Value",
                            "Page", "Source Text"], elements)
    wb.active.title = "File Elements"
    write_sheet(wb.create_sheet("Field Map"), FIELD_MAP_COLS, fmap)
    wb.save(args.output)
    print(f"Wrote {len(elements)} file elements and {len(fmap)} field mappings to {args.output}")


if __name__ == "__main__":
    main()
