"""
Pull every table out of a PDF (default: the HUD HMIS Data Dictionary) into
an Excel workbook.

Usage:
    python scripts/pdf_tables_to_excel.py
    python scripts/pdf_tables_to_excel.py path/or/url/to/file.pdf -o out.xlsx
    python scripts/pdf_tables_to_excel.py --no-merge --pages 10-40

Output workbook:
    Index         one row per table: sheet name, source pages, heading, size
    T001, T002... one sheet per table, preceded by its heading and pages

Tables that run across a page break (same column count, starting at the top
of the next page) are stitched into one table, and a repeated header row on
the continuation page is dropped. Each table is labelled with the nearest
data-element heading above it (e.g. "3.01 Name", "P4 Connection with SOAR"),
carried forward from earlier pages when the page itself has none.
"""
import argparse
import re
import sys
import tempfile
import urllib.request
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_URL = "https://files.hudexchange.info/resources/documents/HMIS-Data-Dictionary.pdf"

# Data-element headings in the HMIS dictionary: "2.02 Project Information",
# "3.917 Prior Living Situation", "4.12 Current Living Situation",
# "P4 Connection with SOAR", "R7 General Health Status", "V2 ...", "W5 ...", "C2 ..."
HEADING_RE = re.compile(r"^(\d+\.\d+[A-Z]?|[A-Z]\d{1,2}[A-Z]?)\s+[A-Z][\w&/(),'\- ]{2,}$")

# Where a table must start / end (as a fraction of page height) to count as
# continuing across a page break.
TOP_FRACTION = 0.20
BOTTOM_FRACTION = 0.80

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def fetch_pdf(source):
    """Return a local path for `source`, downloading it if it is a URL."""
    if not re.match(r"https?://", source):
        return Path(source)
    dest = Path(tempfile.gettempdir()) / Path(source).name
    if not dest.exists():
        print(f"Downloading {source} ...", file=sys.stderr)
        req = urllib.request.Request(source, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req) as resp, open(dest, "wb") as fh:
            fh.write(resp.read())
    return dest


def parse_pages(spec, n_pages):
    """'1-5,8' -> [0, 1, 2, 3, 4, 7] (0-based)."""
    if not spec:
        return list(range(n_pages))
    pages = []
    for part in spec.split(","):
        lo, _, hi = part.partition("-")
        pages.extend(range(int(lo) - 1, int(hi or lo)))
    return [p for p in pages if 0 <= p < n_pages]


def clean_cell(value):
    if value is None:
        return ""
    # Collapse the hard line wraps pdfplumber keeps inside cells.
    return re.sub(r"\s*\n\s*", " ", str(value)).strip()


def clean_rows(rows):
    rows = [[clean_cell(c) for c in row] for row in rows]
    return [r for r in rows if any(r)]


def find_heading(page, table_top):
    """Closest data-element heading above `table_top` on this page, or None."""
    lines = page.extract_text_lines(strip=True) if hasattr(page, "extract_text_lines") else []
    heading = None
    for line in lines:
        if line["bottom"] > table_top:
            break
        text = line["text"].strip()
        if HEADING_RE.match(text):
            heading = text
    return heading


def extract_tables(pdf_path, page_spec=None, merge=True):
    tables = []
    last_heading = ""
    with pdfplumber.open(pdf_path) as pdf:
        for pno in parse_pages(page_spec, len(pdf.pages)):
            page = pdf.pages[pno]
            found = sorted(page.find_tables(), key=lambda t: t.bbox[1])
            if not found:
                # Still track headings so the next table gets the right label.
                h = find_heading(page, page.height)
                last_heading = h or last_heading
                continue
            for tbl in found:
                rows = clean_rows(tbl.extract())
                if not rows:
                    continue
                top, bottom = tbl.bbox[1], tbl.bbox[3]
                heading = find_heading(page, top)
                prev = tables[-1] if tables else None
                is_continuation = (
                    merge
                    and prev is not None
                    and heading is None
                    and prev["last_page"] == pno - 1
                    and prev["ends_low"]
                    and top < page.height * TOP_FRACTION
                    and len(rows[0]) == len(prev["rows"][0])
                )
                if is_continuation:
                    if rows[0] == prev["rows"][0]:
                        rows = rows[1:]  # repeated header
                    prev["rows"].extend(rows)
                    prev["last_page"] = pno
                    prev["ends_low"] = bottom > page.height * BOTTOM_FRACTION
                    continue
                last_heading = heading or last_heading
                tables.append({
                    "heading": last_heading,
                    "first_page": pno,
                    "last_page": pno,
                    "ends_low": bottom > page.height * BOTTOM_FRACTION,
                    "rows": rows,
                })
            # Headings below the last table on the page apply to the next page.
            h = find_heading(page, page.height)
            last_heading = h or last_heading
    return tables


def autosize(ws, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), max_width)


def write_workbook(tables, out_path):
    wb = Workbook()
    index = wb.active
    index.title = "Index"
    index.append(["Sheet", "Heading", "Pages", "Rows", "Columns"])
    for cell in index[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL

    for i, t in enumerate(tables, 1):
        name = f"T{i:03d}"
        pages = (f"{t['first_page'] + 1}" if t["first_page"] == t["last_page"]
                 else f"{t['first_page'] + 1}-{t['last_page'] + 1}")
        n_cols = max(len(r) for r in t["rows"])

        ws = wb.create_sheet(name)
        ws.append([t["heading"] or "(no heading found)"])
        ws.append([f"PDF page(s) {pages}"])
        ws.append([])
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"].font = Font(italic=True)
        for r in t["rows"]:
            ws.append(r)
        for cell in ws[4]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A5"
        autosize(ws)

        index.append([name, t["heading"], pages, len(t["rows"]), n_cols])
        index.cell(row=index.max_row, column=1).hyperlink = f"#'{name}'!A1"
        index.cell(row=index.max_row, column=1).style = "Hyperlink"

    index.freeze_panes = "A2"
    autosize(index, max_width=80)
    wb.save(out_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", nargs="?", default=DEFAULT_URL, help="PDF path or URL")
    ap.add_argument("-o", "--output", default="HMIS-Data-Dictionary-tables.xlsx")
    ap.add_argument("--pages", help="1-based page ranges to scan, e.g. '1-20,35'")
    ap.add_argument("--no-merge", action="store_true", help="don't stitch tables across page breaks")
    args = ap.parse_args(argv)

    pdf_path = fetch_pdf(args.source)
    tables = extract_tables(pdf_path, args.pages, merge=not args.no_merge)
    if not tables:
        sys.exit("No tables found.")
    write_workbook(tables, args.output)
    print(f"Wrote {len(tables)} tables to {args.output}")


if __name__ == "__main__":
    main()
