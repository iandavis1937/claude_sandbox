"""
Pivot the 'All Tables' union sheet written by pdf_tables_to_excel.py into a
table with one row per field of each HMIS data element.

Usage:
    python scripts/union_to_fields.py HMIS-Data-Dictionary-tables.xlsx
    python scripts/union_to_fields.py union.xlsx -o fields.xlsx --sep " | "

Input: the first sheet of the workbook. Columns 1-3 are provenance (Table,
Heading, Pages), column 4 holds the row label ('Element Name', 'Field 1 &
Response', a response code, 'Element Type', ...), column 5 holds its value,
and any further columns hold X marks (e.g. 'Retired', or per-program columns
whose names are given on the element's 'Element Name' row).

Output columns (one row per 'Field N & Response(s)'):
    Element Name, Field Name, Responses   responses collapsed into one cell
    Project Type Applicability
    DE#                                   join key to the HMIS CSV Format
                                          Specifications' DE# column (see below)
    CSV                                   the CSV file(s) the element exports to
    <element attributes>                  Element Type, Funder: Program-Component,
                                          Collection Point, ... (repeated per field)
    <field attributes>                    Dependency Logic, Section, Applies To
    Field Label, Source Table, Pages      where the field came from
    Dependent A, Dependent B, ...         each dependent collapsed into one cell on
                                          the row of the field it depends on: the
                                          rest of its label ("Dependent to Field 6 &
                                          Response 13"), its text, and its responses

PDF quirks handled:
    - a label wrapped over several rows (later rows have no value) is rejoined
    - a value wrapped onto rows with no label is appended to the row above
    - several elements in one table are split at each 'Element Name' row

Join key (DE#): <element number>.<field number>, e.g. "2.02.6" for 2.02
Project Information Field 6 (Project Type) or "R13.9" for R13 Family Critical
Issues Field 9 - the same numbering the CSV spec's DE# column uses. Element
numbers come from the "2.02 Project Information"-style headings in the input's
Heading column (matched to the element by name, since that column can sit on
the wrong table), with ELEMENT_NUMBERS supplying the elements no heading names.
KEY_OVERRIDES holds the few places the spec numbers a field differently (e.g.
"3.08" rather than "3.08.1"). Dependents are keyed <element>.<letter> in the
CSV spec (e.g. "2.06.A"); they live in the Dependent columns here, not in DE#.
Fields stored in the spec's long-format files (Disabilities.csv, Services.csv,
the InformationDate/DateCreated columns, ...) have no DE# there, so DE# cannot
join them.

List-like values (response options, project types, funders, dependency
conditions) are separated with --sep, ";" plus a newline by default. A wrapped
line counts as a new list item when it starts like one ("3: PH – ...",
"HUD: CoC ...", "Pay for Success – ...", "[If Yes ..."); otherwise it is prose
and is joined with a space.
"""
import argparse
import re
import sys
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")

# Response codes: "1", "-1", "102-112", "1…n".
CODE_RE = re.compile(r"^-?\d+(\s*[-–…]+\s*(\d+|n))?$")

# Rows that start a field: "Field 2 & Responses", "Dependent A – Dependent to ...",
# "Dependency A, dependent to Field 1 Response 1".
FIELD_RE = re.compile(r"^(Field\b|Dependent\b|Dependency [A-Z]\b)")

# A dependent's heading: letter, then the rest of the label.
DEP_RE = re.compile(r"^Dependen\w*\s+([A-Z])\b[\s,:–-]*(.*)$")
# A field's own heading, as opposed to a reference like "Field 6 & Response 6 or".
FIELD_HEAD_RE = re.compile(r"^Field\s*\d+[A-Z]?(\s*&\s*Responses?)?(\s*$|\s+[–-])")
DEP_START_RE = re.compile(r"Dependen\w* [A-Z]\b\s*([–,-]|to\b)")

# Where a dependent points: "to Field 6 ..." or "to Dependent B ...".
DEP_REF_RE = re.compile(r"Fields?\s*#?\s*(\d+)|Dependen\w*\s+([A-Z])\b")

# Labels that mark the start of the element-level attribute block. Anything
# unrecognised after one of these is also element-level; before, field-level.
ELEMENT_ATTR_KEYS = {
    "elementtype", "funderprogramcomponent", "projecttypeapplicability",
    "datacollectedabout", "collectionpoint", "systemlogicothersystemissues",
    "csv", "relationshiptoenrollmentidanenrollment", "relationshiptopersonalidaclient",
}

# Preferred spellings for labels whose wrapping varies between tables, keyed by
# the prefix of key(label) they cover (so 'System Logic & Other System Issues
# Other System Issues' still maps to one column).
CANONICAL = {
    "funderprogramcomponent": "Funder: Program-Component",
    "systemlogicothersystemissues": "System Logic & Other System Issues",
    "relationshiptoenrollmentidanenrollment": "Relationship to Enrollment ID (an enrollment)",
    "relationshiptopersonalidaclient": "Relationship to Personal ID (a client)",
}


# "2.02 Project Information", "3.917A Prior Living Situation", "R15 Commercial ..."
ELEMENT_HEADING_RE = re.compile(r"^(\d+\.\d+|[A-Z]\d{1,2})[A-Z]?\s+(.+)$")

# Element numbers for elements the input's headings don't name. Sources: the
# HMIS CSV Format Specifications text (e.g. "Project.csv includes ... 2.02
# Project Information", "V7 SSVF HP Targeting Criteria") except where noted.
ELEMENT_NUMBERS = {
    "Project Information": "2.02",
    "Name": "3.01",
    "Project Start Date": "3.10",
    "Housing Move-In Date": "3.20",
    "Chronic Health Condition": "4.07",
    "Substance Use Disorder": "4.10",
    "Coordinated Entry Assessment": "4.19",
    "Coordinated Entry Event": "4.20",
    "Project Identifier": "5.05",         # not in the CSV spec; follows 5.04 in the dictionary
    "Personal Identifier": "5.08",
    "Implementation Identifier": "5.10",  # not in the CSV spec; follows 5.09 in the dictionary
    "Services Provided – PATH Funded": "P1",
    "Referrals Provided PATH": "P2",
    "RHY – BCP Status": "R2",
    "RHY Service Connections": "R14",
    "Safe and Appropriate Exit": "R19",
    "Veterans Information": "V1",
    "Services Provided – SSVF": "V2",
    "Financial Assistance – SSVF": "V3",
    "Percent of AMI (SSVF Eligibility)": "V4",
    "HP Targeting Criteria": "V7",
    "Services Provided – HOPWA": "W1",
    "Financial Assistance – HOPWA": "W2",
    "Medical Assistance": "W3",
    "T-Cell (CD4) and Viral Load": "W4",
}

# "Field 6 & Responses" -> "6", "Field 2A – for Institutional Situations" -> "2A"
FIELD_NUM_RE = re.compile(r"^Field\s*#?\s*(\d+[A-Z]?)\b")

# Where the CSV spec's DE# differs from <element>.<field>.
KEY_OVERRIDES = {
    "3.08.1": "3.08",    # Enrollment.csv DisablingCondition: single-field element, no ".1"
    "4.21.1": "4.21",    # Client.csv Sex
    "V10.1": "V10",      # Enrollment.csv MentalHealthConsultation
    "5.07.1": "5.07.01", # User.csv UserID is printed "5.07.01"
    "W2.3": "W2",        # Services.csv FAAmount is keyed "W2 & V3.3"
    "W3.2": "W3.3",      # IncomeBenefits.csv ADAP: the dictionary numbers it Field 2, the spec W3.3
}

DEFAULT_SEP = ";\n"

# A wrapped value line that starts a new list item rather than continuing prose.
LIST_ITEM_RE = re.compile(r"^(\d{1,3}: |[A-Z]{2,}: |[A-Z][\w/&]*( [\w/&]+){0,3} – |\[If\b)")

# "0: Emergency Shelter 1: Day Shelter" packed into one cell.
INLINE_ITEM_RE = re.compile(r" (?=\d{1,3}: [A-Z])")


def text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


def join(a, b):
    """Join wrapped text, without a space after a line-ending hyphen ('Program-')."""
    if not a or not b:
        return a or b
    return a + b if re.search(r"\S[-–]$", a) else f"{a} {b}"


def join_item(a, b, sep):
    """Join a wrapped value line: `sep` if it starts a new list item, else as prose."""
    return f"{a}{sep}{b}" if a and b and LIST_ITEM_RE.match(b) else join(a, b)


def split_inline(value, sep):
    """Separate a numbered list run together in one cell, skipping '(subtype 2: ...)'."""
    if not re.match(r"\d{1,3}: ", value):
        return value
    out, last = [], 0
    for m in INLINE_ITEM_RE.finditer(value):
        if value[last:m.start()].count("(") == value[last:m.start()].count(")"):
            out.append(value[last:m.start()])
            last = m.end()
    return sep.join(out + [value[last:]])


def key(label):
    return re.sub(r"[^a-z0-9]", "", label.lower())


def attr_name(label):
    k = key(label)
    return next((name for c, name in CANONICAL.items() if k.startswith(c)), label)


def is_element_attr(k):
    return k in ELEMENT_ATTR_KEYS or bool(re.fullmatch(r"\d{4}revisionsummary", k))


def is_section(label):
    # All-caps divider such as "DEPENDENCIES"; short tokens like "CSV" or "ID" are not.
    return len(key(label)) > 3 and bool(re.search(r"[A-Z]{2}", label)) and label == label.upper()


def starts_label(label, value, prev_label):
    """True if `label` begins a new row rather than continuing a wrapped one."""
    if CODE_RE.match(label):
        return bool(value)  # a bare number with no value is the tail of a wrapped label
    if FIELD_RE.match(label):
        # 'Dependent B – Dependent to' / 'Field 6 & Response 6 or' / 'Dependent A = 1':
        # the Field/Dependent lines after a dependent heading are its wrapped label.
        # A label ending in "to"/"or"/"&" continues even when the next line has text.
        dangling = re.search(r"(\b(to|or|and)|[&–-])$", prev_label)
        wrapped = (re.match(r"Dependen", prev_label) and not DEP_START_RE.match(label)
                   and (dangling or (not value and not FIELD_HEAD_RE.match(label))))
        return not wrapped
    if re.match(r"(Element Name|Dependency Logic|Header:)", label):
        return True
    if is_section(label):
        return bool(value) or not is_section(prev_label)
    k = key(label)
    return len(k) >= 3 and (any(a.startswith(k) for a in ELEMENT_ATTR_KEYS)
                            or bool(re.match(r"\d{4}revision", k)))


def logical_rows(sheet_rows, sep=DEFAULT_SEP):
    """Rejoin wrapped labels/values. Yields dicts: table, pages, label, value, cells."""
    cur = None
    for r in sheet_rows:
        label, value = text(r[3]), split_inline(text(r[4]), sep)
        cells = [text(v) for v in r[5:]]
        if not label and not value and not any(cells):
            continue
        same_table = cur is not None and r[0] == cur["table"]
        if same_table and (not label or not starts_label(label, value, cur["label"])):
            cur["label"] = join(cur["label"], label)  # wrapped label and/or value
            cur["value"] = join_item(cur["value"], value, sep)
            continue
        if cur:
            yield cur
        cur = {"table": r[0], "pages": text(r[2]), "label": label, "value": value, "cells": cells,
               "first_label": label}
    if cur:
        yield cur


def parse_elements(rows, extra_names, sep=DEFAULT_SEP):
    """Group logical rows into elements -> fields."""
    elements = []
    el = field = None
    section = ""
    for row in logical_rows(rows, sep):
        label, value, k = row["label"], row["value"], key(row["label"])

        if k == "elementname":
            # Per-program column names may be given on the Element Name row.
            sub = {j: c for j, c in enumerate(row["cells"]) if c}
            el = {"name": value, "attrs": OrderedDict(), "fields": [], "deps": [], "in_attrs": False,
                  "table": row["table"], "pages": row["pages"], "subheaders": sub}
            elements.append(el)
            field, section = None, ""
            continue
        if el is None:
            continue  # rows before the first Element Name

        def marks():
            return [el["subheaders"].get(j, extra_names[j])
                    for j, c in enumerate(row["cells"]) if c]

        if CODE_RE.match(row["first_label"]) and not el["in_attrs"]:
            if field is None:  # responses with no field row: attach to a blank field
                field = new_field(el, "", "", section, row)
            code = row["first_label"]
            resp = join(f"{code}: {value}", label[len(code):].strip())
            if marks():
                resp += f" [{', '.join(marks())}]"
            field["responses"].append(resp)
        elif FIELD_RE.match(label) and not el["in_attrs"]:
            dep = DEP_RE.match(label)
            if dep:
                field = {"letter": dep.group(1), "tail": dep.group(2), "name": value,
                         "responses": [], "attrs": OrderedDict(), "prev": el["fields"][-1:]}
                el["deps"].append(field)
            else:
                field = new_field(el, label, value, section, row)
            if marks():
                field["attrs"]["Applies To"] = sep.join(marks())
        elif k == "header" and field is not None:
            field["responses"].append(f"[{value}]")  # group heading within a response list
        elif is_element_attr(k) or el["in_attrs"]:
            el["in_attrs"] = True
            name = attr_name(label)
            el["attrs"][name] = join(el["attrs"].get(name, ""), value) if name in el["attrs"] else value
        elif is_section(label):
            section = join(label, value)  # e.g. "DEPENDENCIES FOR INSTITUTIONAL SITUATIONS"
        elif field is not None:
            field["attrs"][label] = join(field["attrs"].get(label, ""), value)
        else:
            el["attrs"][label] = value
    for el in elements:
        attach_dependents(el, sep)
    return elements


def attach_dependents(el, sep):
    """Collapse each dependent into a 'Dependent X' cell on the field it depends on."""
    by_num = {}
    for f in el["fields"]:
        m = re.match(r"Field\s*#?\s*(\d+)(?!\w)", f["label"])
        if m:
            by_num.setdefault(m.group(1), f)
    by_letter = {d["letter"]: d for d in el["deps"]}

    def parent(d, seen=()):
        m = DEP_REF_RE.search(d["tail"])
        if m and m.group(1) in by_num:
            return by_num[m.group(1)]
        if m and m.group(2) in by_letter and m.group(2) not in seen:
            return parent(by_letter[m.group(2)], (*seen, d["letter"]))
        return d["prev"][0] if d["prev"] else None  # field just above the dependent

    for d in el["deps"]:
        target = parent(d)
        if target is None:  # dependents with no field at all: give them a blank row
            target = new_field(el, "", "", "", {"table": el["table"]})
        attrs = [f"{k}: {v}" for k, v in d["attrs"].items()]
        cell = sep.join(x for x in [d["tail"], d["name"], *d["responses"], *attrs] if x)
        col = f"Dependent {d['letter']}"
        deps = target.setdefault("deps", {})
        deps[col] = f"{deps[col]}{sep}{cell}" if col in deps else cell


def new_field(el, label, name, section, row):
    f = {"label": label, "name": name, "responses": [], "attrs": OrderedDict(),
         "table": row["table"], "pages": el["pages"]}
    if section:
        f["attrs"]["Section"] = section
    el["fields"].append(f)
    return f


def name_key(name):
    """Element name for matching: lower-case, no parentheticals or punctuation."""
    return re.sub(r"[^a-z0-9]", "", re.sub(r"\(.*?\)", "", name.lower()))


def harvest_element_numbers(sheet_rows):
    """{name_key: element number} from 'N.NN Element Name' headings in column 2."""
    found = {}
    for r in sheet_rows:
        m = ELEMENT_HEADING_RE.match(text(r[1]))
        if m:
            found.setdefault(name_key(m.group(2)), m.group(1))
    return found


def element_number(name, harvested):
    overrides = {name_key(k): v for k, v in ELEMENT_NUMBERS.items()}
    k = name_key(name)
    if k in overrides:
        return overrides[k]
    if k in harvested:
        return harvested[k]
    # Headings can run on into following text ("5.04 Information Date Corrected ...").
    return next((num for hk, num in harvested.items() if k and hk.startswith(k)), "")


def de_number(el_num, field_num):
    key = f"{el_num}.{field_num}"
    return KEY_OVERRIDES.get(key, key)


def to_records(elements, sep, harvested=None):
    records = []
    for el in elements:
        el_num = element_number(el["name"], harvested or {})
        if not el_num:
            print(f"Warning: no element number for '{el['name']}'; add it to ELEMENT_NUMBERS",
                  file=sys.stderr)
        for f in el["fields"] or [{"label": "", "name": "", "responses": [], "attrs": {},
                                   "table": el["table"], "pages": el["pages"]}]:
            m = FIELD_NUM_RE.match(f["label"])
            records.append({
                "Element Name": el["name"],
                "Field Name": f["name"],
                "Responses": sep.join(f["responses"]),
                "DE#": de_number(el_num, m.group(1)) if el_num and m else "",
                "_deps": f.get("deps", {}),
                "_element": el["attrs"],
                "_field": f["attrs"],
                "Field Label": f["label"],
                "Source Table": f["table"],
                "Pages": f["pages"],
            })
    return records


def write_output(records, out_path):
    el_cols, f_cols = [], []
    dep_cols = sorted({c for rec in records for c in rec["_deps"]})
    for rec in records:
        el_cols += [c for c in rec["_element"] if c not in el_cols]
        f_cols += [c for c in rec["_field"] if c not in f_cols]
    lead = ["Element Name", "Field Name", "Responses", "Project Type Applicability", "DE#", "CSV"]
    columns = (lead + [c for c in el_cols + f_cols if c not in lead]
               + ["Field Label", "Source Table", "Pages"] + dep_cols)

    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    ws.append(columns)
    for rec in records:
        merged = {**rec, **rec["_deps"], **rec["_element"], **rec["_field"]}
        ws.append([merged.get(c, "") for c in columns])

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 45 if c in ("Responses", *dep_cols, *el_cols) else 28
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)
    return columns


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", help="workbook whose first sheet is the 'All Tables' union")
    ap.add_argument("-o", "--output", help="default: <input>_by_field.xlsx")
    ap.add_argument("--sep", default=DEFAULT_SEP,
                    help="separator between list items such as responses and project types "
                         "(default: ';' + newline; escapes like '\\n' are allowed)")
    args = ap.parse_args(argv)

    src = Path(args.input)
    out = Path(args.output) if args.output else src.with_name(f"{src.stem}_by_field.xlsx")
    ws = load_workbook(src, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    extra_names = [text(h) for h in header[5:]]

    sep = args.sep.encode().decode("unicode_escape")
    rows = list(rows)
    elements = parse_elements(rows, extra_names, sep)
    records = to_records(elements, sep, harvest_element_numbers(rows))
    columns = write_output(records, out)
    print(f"Wrote {len(records)} fields from {len(elements)} elements to {out}")
    print(f"Columns: {', '.join(columns)}")


if __name__ == "__main__":
    main()
