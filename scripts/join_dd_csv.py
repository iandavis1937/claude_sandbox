"""
Join the field-level HMIS Data Dictionary table (union_to_fields.py output) to
the HMIS CSV Format Specifications field list (pdf_tables_to_excel.py output)
on DE#.

Usage:
    python scripts/join_dd_csv.py HUD_field_reqs_by_project_type_by_field.xlsx HMIS_CSV_fields.xlsx
    python scripts/join_dd_csv.py dd_fields.xlsx csv_fields.xlsx -o joined.xlsx
    python scripts/join_dd_csv.py dd_fields.xlsx csv_fields.xlsx --crosswalk HMIS_CSV_crosswalk.xlsx

Inputs:
    dictionary  first sheet; one row per field with a 'DE#' column and
                'Dependent A', 'Dependent B', ... columns
    CSV spec    first sheet ('All Tables'); columns Heading (the CSV file),
                DE#, Name, Type, List, Null, Notes
    crosswalk   optional; the 'Field Map' sheet from csv_spec_crosswalk.py

Matching (a full outer join, one output row per matched pair):
    field       CSV DE# equals the dictionary row's DE#, e.g. 2.02.6
    dependent   CSV DE# is <element>.<letter> (2.06.A) or <element>.<field><letter>
                (4.04.3A) and that element has a 'Dependent <letter>' cell
    crosswalk   a crosswalk Field Map row links the CSV column to the field, under
                the given Condition (e.g. Disabilities.csv DisabilityType <-> 4.05.2
                Physical Disability when DisabilityType = 5). Element-level rows
                resolve by role: 'date' -> the element's first field named
                "...date...", 'value'/'discriminator' -> its first other field.
    CSV only    no dictionary match (incl. the spec's DE#-less columns such as
                DateCreated, and long-format files like Disabilities.csv)
    dictionary only
                no CSV row carries this DE#

A bare element number ("5.08") matches that element's Field 1 when the element
has only one field. A CSV DE# naming several elements ("V2.D & V8.A", "3.917.2A & 2B",
"W5.A & .B", "V2.A V2.B V2.C") is split, and the CSV row matches each part.
"""
import argparse
import re
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")

# Element number at the start of a DE#: "2.02", "3.917", "R13", "V7".
ELEMENT_RE = re.compile(r"^(\d+\.\d+|[A-Z]\d{1,2})")
# A dependent's DE#: "2.06.A", "V7.C", "4.04.3A" (element, optional field, letter).
DEP_KEY_RE = re.compile(r"^(\d+\.\d+|[A-Z]\d{1,2})\.\d*([A-Z])$")

CSV_COLS = ["CSV File", "CSV Field", "Type", "List", "Null", "Notes"]


def read_sheet(path):
    ws = load_workbook(path, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    return header, [dict(zip(header, r)) for r in rows if any(v not in (None, "") for v in r)]


def split_de(de):
    """'V2.D & V8.A' -> ['V2.D', 'V8.A']; '3.917.2A & 2B' -> ['3.917.2A', '3.917.2B']."""
    keys, prev = [], ""
    for tok in re.split(r"\s*&\s*|\s+", str(de or "").strip()):
        if not tok or tok == "*":
            continue
        if tok.startswith(".") and prev:                       # "W5.A & .B"
            tok = prev.rsplit(".", 1)[0] + tok
        elif re.fullmatch(r"\d+[A-Z]?", tok) and prev:         # "3.917.2A & 2B"
            tok = prev.rsplit(".", 1)[0] + "." + tok
        keys.append(tok)
        prev = tok
    return keys


def read_crosswalk(path):
    ws = load_workbook(path, read_only=True)["Field Map"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    index = defaultdict(list)
    for r in rows:
        e = dict(zip(header, r))
        index[(e["CSV File"], e["CSV Field"])].append(e)
    return index


def pick_field(rows, role):
    """The element's field for a crosswalk role: its date field, or its first other field."""
    is_date = [bool(re.search(r"\bdate\b", str(r.get("Field Name") or ""), re.I)) for r in rows]
    want = [r for r, d in zip(rows, is_date) if d == (role == "date")]
    if role == "date":
        return want[:1]  # no date field (e.g. P4): no match
    return (want or rows)[:1]


def join(dd_header, dd_rows, csv_rows, crosswalk=None):
    by_key = defaultdict(list)
    deps = defaultdict(list)  # (element, letter) -> [(dd row, dependent cell)]
    dep_cols = [c for c in dd_header if re.fullmatch(r"Dependent [A-Z]", c)]
    for r in dd_rows:
        key = str(r.get("DE#") or "")
        if key:
            by_key[key].append(r)
        el = ELEMENT_RE.match(key)
        for c in dep_cols:
            if el and r.get(c):
                deps[(el.group(1), c[-1])].append((r, r[c]))

    # Distinct field keys per element, for CSV DE#s that give only the element.
    fields_per_element = defaultdict(set)
    element_rows = defaultdict(list)
    for key, rows in by_key.items():
        el = ELEMENT_RE.match(key)
        if el:
            fields_per_element[el.group(1)].add(key)
            element_rows[el.group(1)] += rows

    def resolve(entry):
        """[(dictionary row, dependent cell or None)] for a crosswalk Field Map row."""
        key = str(entry["DE#"])
        if entry["Level"] == "element":
            return [(r, None) for r in pick_field(element_rows.get(key, []), entry["Role"])]
        if key not in by_key and fields_per_element.get(key) == {f"{key}.1"}:
            key = f"{key}.1"
        if key in by_key:
            return [(r, None) for r in by_key[key]]
        dep = DEP_KEY_RE.match(key)
        return deps.get((dep.group(1), dep.group(2)), []) if dep else []

    out, used = [], set()
    for c in csv_rows:
        csv_part = {"CSV File": c.get("Heading"), "CSV Field": c.get("Name"),
                    **{k: c.get(k) for k in ("Type", "List", "Null", "Notes")},
                    "CSV DE#": c.get("DE#")}
        matched, paired = False, set()  # paired: dictionary rows already joined to c
        for key in split_de(c.get("DE#")):
            # Client.csv's PersonalID is "5.08": a bare element number is its Field 1,
            # but only when the element has a single field (P2 alone is not P2.1).
            if key not in by_key and fields_per_element.get(key) == {f"{key}.1"}:
                key = f"{key}.1"
            if key in by_key:
                for r in by_key[key]:
                    out.append({**r, **csv_part, "Match": "field", "DE#": key})
                    used.add(id(r))
                    paired.add((id(r), None))
                matched = True
                continue
            dep = DEP_KEY_RE.match(key)
            if dep and (dep.group(1), dep.group(2)) in deps:
                for r, cell in deps[(dep.group(1), dep.group(2))]:
                    out.append({**r, **csv_part, "Match": "dependent", "DE#": key,
                                "Matched Dependent": cell})
                matched = True
        for entry in (crosswalk or {}).get((c.get("Heading"), c.get("Name")), []):
            for r, dep_cell in resolve(entry):
                if (id(r), dep_cell) in paired:
                    continue
                paired.add((id(r), dep_cell))
                out.append({**r, **csv_part, "Match": "crosswalk",
                            "DE#": entry["DE#"] if dep_cell else r.get("DE#"),
                            "Matched Dependent": dep_cell,
                            "Condition": entry["Condition"], "Crosswalk Rule": entry["Rule"]})
                if not dep_cell:
                    used.add(id(r))
                matched = True
        if not matched:
            out.append({"Match": "CSV only", "DE#": c.get("DE#"), **csv_part})
    for r in dd_rows:
        if id(r) not in used and r.get("DE#"):
            out.append({"Match": "dictionary only", **r})
    return out


def write(rows, dd_header, out_path):
    dd_cols = [c for c in dd_header if c != "DE#"]
    columns = (["Match", "DE#"] + CSV_COLS + ["CSV DE#", "Matched Dependent", "Condition", "Crosswalk Rule"]
               + dd_cols)

    wb = Workbook()
    ws = wb.active
    ws.title = "Joined"
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 12 if c in ("Match", "Type", "List", "Null") else 28
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("dictionary", help="union_to_fields.py output")
    ap.add_argument("csv_spec", help="pdf_tables_to_excel.py output for the CSV spec")
    ap.add_argument("-o", "--output", default="HMIS_DD_CSV_join.xlsx")
    ap.add_argument("--crosswalk", help="csv_spec_crosswalk.py output")
    args = ap.parse_args(argv)

    dd_header, dd_rows = read_sheet(args.dictionary)
    _, csv_rows = read_sheet(args.csv_spec)
    rows = join(dd_header, dd_rows, csv_rows, read_crosswalk(args.crosswalk) if args.crosswalk else None)
    write(rows, dd_header, args.output)

    counts = defaultdict(int)
    for r in rows:
        counts[r["Match"]] += 1
    print(f"Wrote {len(rows)} rows to {args.output}: "
          + ", ".join(f"{n} {m}" for m, n in sorted(counts.items(), key=lambda x: -x[1])))


if __name__ == "__main__":
    main()
